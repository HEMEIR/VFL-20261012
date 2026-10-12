import sys
from docx import Document

p = r'C:\Users\HEMEI\Desktop\新建 DOCX 文档.docx'
d = Document(p)
out = []
out.append('段落数: %d  表格数: %d' % (len(d.paragraphs), len(d.tables)))
out.append('=' * 72)
for i, para in enumerate(d.paragraphs):
    t = para.text.rstrip()
    if t:
        out.append('[%3d][%s] %s' % (i, para.style.name, t))
for ti, tb in enumerate(d.tables):
    out.append('--- 表格 %d (%d行 x %d列) ---' % (ti, len(tb.rows), len(tb.columns)))
    for r in tb.rows:
        out.append('  | ' + ' | '.join(c.text.strip().replace('\n', ' ') for c in r.cells))

with open(r'D:\MakeCode\VisualStudio\WorkPlace\VFL\VFL-Security-Challenge-v1.0.7\outputs\_docx_dump.txt', 'w', encoding='utf-8') as fh:
    fh.write('\n'.join(out))
print('written', len(out), 'lines')
